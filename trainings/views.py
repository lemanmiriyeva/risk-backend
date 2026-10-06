import logging
from io import BytesIO

from django.db import transaction
from django.db.models import Avg, Count, Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework import filters, mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from authentication.models import User
from authentication.permissions import HasSystemAccess
from core.permissions import ModuleAccessPermission
from .models import QuizAnswer, QuizAttempt, Training, TrainingFeedback, TrainingProgress
from .permissions import (
    TRAININGS_MODULE_CODE,
    TrainingMaterialsPermission,
    TrainingStatisticsPermission,
    can_manage_materials,
    can_manage_statistics,
    can_view_materials,
    can_view_statistics,
    materials_audience_user_ids,
    org_scope_q,
    statistics_user_scope,
)
from .serializers import (
    QuizAttemptDetailSerializer,
    QuizAttemptSerializer,
    QuizQuestionManageSerializer,
    QuizQuestionPublicSerializer,
    QuizReplaceSerializer,
    QuizSubmitSerializer,
    TrainingFeedbackSerializer,
    TrainingFeedbackStatSerializer,
    TrainingProgressStatSerializer,
    TrainingSerializer,
    _user_info,
)

logger = logging.getLogger("colored")

# Baxış yoxlaması üçün tolerantlıq parametrləri (bax: _apply_position).
# Brauzer ~5 saniyədən bir irəliləyiş göndərir; RATE oynatma sürətinə və
# şəbəkə gecikməsinə imkan verir, amma videonu irəli "atlamağa" imkan vermir.
HEARTBEAT_TOLERANCE_SECONDS = 5
MAX_PLAYBACK_RATE = 1.5
END_TOLERANCE_SECONDS = 3

TRAINING_LINK = "/telimler/materiallar/{id}"


def _apply_position(progress, position, now, duration):
    """
    Brauzerin bildirdiyi mövqeyi server vaxtı ilə məhdudlaşdırır: son
    aktivlikdən bəri real keçən vaxtdan çox irəliləmək mümkün deyil.
    """
    try:
        position = float(position)
    except (TypeError, ValueError):
        raise ValidationError({"position": "Yanlış dəyər."})

    elapsed = 0.0
    if progress.last_heartbeat_at:
        elapsed = max((now - progress.last_heartbeat_at).total_seconds(), 0.0)
    allowed = progress.max_position + elapsed * MAX_PLAYBACK_RATE + HEARTBEAT_TOLERANCE_SECONDS
    position = max(0.0, min(position, allowed))
    if duration:
        position = min(position, float(duration))
    if position > progress.max_position:
        progress.max_position = position
    progress.last_heartbeat_at = now


def _required_end_position(duration):
    return max(duration - END_TOLERANCE_SECONDS, duration * 0.95)


def _progress_payload(progress, training):
    if not progress:
        return {"status": None, "max_position": 0, "duration_seconds": training.duration_seconds}
    return {
        "status": progress.status,
        "max_position": round(progress.max_position, 1),
        "duration_seconds": training.duration_seconds,
        "attempts_count": progress.attempts_count,
        "completed_at": progress.completed_at,
    }


def _date_range_filter(qs, request, field):
    date_from = parse_date(request.query_params.get("date_from") or "")
    date_to = parse_date(request.query_params.get("date_to") or "")
    if date_from:
        qs = qs.filter(**{f"{field}__date__gte": date_from})
    if date_to:
        qs = qs.filter(**{f"{field}__date__lte": date_to})
    return qs


USER_SEARCH_FIELDS = [
    "user__firstname", "user__lastname", "user__username", "user__department__title",
]


class StatsPagination(PageNumberPagination):
    page_size = 20
    page_size_query_param = "page_size"
    max_page_size = 500


# --------------------------------------------------------------------------- #
#  Təlim materialları                                                         #
# --------------------------------------------------------------------------- #

class TrainingViewSet(viewsets.ModelViewSet):
    """
    Təlim materialları (videolar).

    - Siyahı/baxış: alt moduluna girişi olan hər kəs (yalnız aktiv təlimlər).
    - Yaratma/redaktə/silmə və quiz tərtibi: yalnız "Təlim materialları" admini.
    - start / progress / complete: videoya baxışın izlənməsi.
    - quiz (GET) + quiz/submit: video SONA QƏDƏR baxıldıqdan sonra açılır.
    - feedback: istifadəçinin öz rəyi.
    """

    serializer_class = TrainingSerializer
    permission_classes = [TrainingMaterialsPermission]
    pagination_class = None
    filter_backends = [filters.SearchFilter]
    search_fields = ["title", "description"]
    user_actions = ("start", "progress", "complete", "quiz_submit", "feedback")

    def get_queryset(self):
        user = self.request.user
        qs = (
            Training.objects.select_related("organization", "created_by")
            .annotate(questions_total=Count("quiz_questions", distinct=True))
            .filter(org_scope_q(user))
        )
        # Admin arxivləşdirilmiş (deaktiv) təlimləri də görür ki, yenidən aktiv edə bilsin.
        if not can_manage_materials(user):
            qs = qs.filter(is_active=True)
        return qs.order_by("-created_at")

    def get_serializer_context(self):
        context = super().get_serializer_context()
        if self.action == "list":
            user = self.request.user
            context["progress_map"] = {
                p.training_id: p for p in TrainingProgress.objects.filter(user=user)
            }
            attempts_map = {}
            for a in QuizAttempt.objects.filter(user=user):
                attempts_map.setdefault(a.training_id, []).append(a)
            context["attempts_map"] = attempts_map
            context["feedback_map"] = {
                f.training_id: f for f in TrainingFeedback.objects.filter(user=user)
            }
        return context

    # ---- CRUD ---------------------------------------------------------------

    def _resolve_organization(self, serializer):
        user = self.request.user
        organization = serializer.validated_data.get("organization")
        # Superuser olmayan admin yalnız ÖZ qurumu üçün (və ya hamıya aid) yükləyə bilər.
        if not user.is_superuser and organization is not None:
            if organization.id != getattr(user, "organization_id", None):
                organization = user.organization if user.organization_id else None
        return organization

    def perform_create(self, serializer):
        training = serializer.save(
            created_by=self.request.user, organization=self._resolve_organization(serializer),
        )
        logger.info(f"{self.request.user} yeni təlim yüklədi: {training}")
        if training.is_active:
            self._notify_audience(training)

    def perform_update(self, serializer):
        kwargs = {}
        if "organization" in serializer.validated_data:
            kwargs["organization"] = self._resolve_organization(serializer)
        was_active = serializer.instance.is_active
        training = serializer.save(**kwargs)
        if training.is_active and not was_active:
            self._notify_audience(training)

    def _notify_audience(self, training):
        try:
            from notifications.services import notify_many

            recipient_ids = materials_audience_user_ids(training) - {self.request.user.id}
            notify_many(
                User.objects.filter(id__in=recipient_ids),
                title="Yeni təlim materialı",
                body=training.title,
                link=TRAINING_LINK.format(id=training.id),
                related_app="trainings",
                related_object_id=training.id,
            )
        except Exception as exc:  # bildiriş xətası təlimin yüklənməsinə mane olmamalıdır
            logger.error(f"Təlim bildirişi göndərilmədi: {exc}")

    def destroy(self, request, *args, **kwargs):
        training = self.get_object()
        has_history = (
            training.progress_records.exists()
            or training.quiz_attempts.exists()
            or training.feedback.exists()
        )
        if has_history:
            # Statistika itməsin deyə tarixçəsi olan təlim silinmir, arxivləşdirilir.
            training.is_active = False
            training.save(update_fields=["is_active", "updated_at"])
            return Response(
                {"detail": "Bu təlimin statistikası olduğu üçün silinmədi, arxivləşdirildi (deaktiv edildi).",
                 "archived": True},
                status=status.HTTP_200_OK,
            )
        training.video.delete(save=False)
        if training.thumbnail:
            training.thumbnail.delete(save=False)
        training.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

    # ---- Baxışın izlənməsi --------------------------------------------------

    def _get_active_training(self):
        training = self.get_object()
        if not training.is_active:
            raise ValidationError("Bu təlim arxivləşdirilib.")
        return training

    @action(detail=True, methods=["post"])
    def start(self, request, pk=None):
        """
        Video səhifəsi açılanda çağırılır. Video hələ sona qədər baxılmayıbsa,
        irəliləyiş SIFIRLANIR - yarımçıq baxış sayılmır, video əvvəldən izlənməlidir.
        """
        training = self._get_active_training()
        now = timezone.now()

        duration = request.data.get("duration")
        if not training.duration_seconds and duration:
            try:
                training.duration_seconds = int(round(float(duration)))
                training.save(update_fields=["duration_seconds"])
            except (TypeError, ValueError):
                pass

        with transaction.atomic():
            progress, created = TrainingProgress.objects.select_for_update().get_or_create(
                user=request.user, training=training,
                defaults={"first_started_at": now},
            )
            if not progress.is_completed:
                progress.max_position = 0
                progress.session_started_at = now
                progress.last_heartbeat_at = now
                progress.attempts_count += 1
                progress.save()
        return Response(_progress_payload(progress, training))

    @action(detail=True, methods=["post"])
    def progress(self, request, pk=None):
        training = self._get_active_training()
        progress = TrainingProgress.objects.filter(user=request.user, training=training).first()
        if not progress or not progress.session_started_at:
            raise ValidationError("Baxış sessiyası başlanmayıb.")
        if not progress.is_completed:
            _apply_position(progress, request.data.get("position"), timezone.now(), training.duration_seconds)
            progress.save(update_fields=["max_position", "last_heartbeat_at"])
        return Response(_progress_payload(progress, training))

    @action(detail=True, methods=["post"])
    def complete(self, request, pk=None):
        """Video sona çatanda çağırılır. Server baxışın həqiqətən sona çatdığını yoxlayır."""
        training = self._get_active_training()
        progress = TrainingProgress.objects.filter(user=request.user, training=training).first()
        if not progress or not progress.session_started_at:
            raise ValidationError("Baxış sessiyası başlanmayıb.")
        if progress.is_completed:
            return Response(_progress_payload(progress, training))

        if not training.duration_seconds:
            raise ValidationError("Videonun müddəti məlum deyil.")

        now = timezone.now()
        _apply_position(progress, request.data.get("position"), now, training.duration_seconds)
        if progress.max_position < _required_end_position(training.duration_seconds):
            progress.save(update_fields=["max_position", "last_heartbeat_at"])
            raise ValidationError("Video sona qədər izlənilməyib.")

        progress.status = TrainingProgress.STATUS_COMPLETED
        progress.completed_at = now
        progress.save()
        logger.info(f"{request.user} təlimi tam izlədi: {training}")
        return Response(_progress_payload(progress, training))

    # ---- Quiz ----------------------------------------------------------------

    @action(detail=True, methods=["get", "put"])
    def quiz(self, request, pk=None):
        training = self.get_object()
        is_manager = can_manage_materials(request.user)

        if request.method == "PUT":
            # İcazə artıq TrainingMaterialsPermission-da yoxlanılıb (yalnız admin).
            serializer = QuizReplaceSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            serializer.save(training)
            logger.info(f"{request.user} təlim quiz-ini yenilədi: {training}")
            questions = training.quiz_questions.prefetch_related("options")
            return Response({"questions": QuizQuestionManageSerializer(questions, many=True).data})

        questions = training.quiz_questions.prefetch_related("options")
        if is_manager and request.query_params.get("manage"):
            return Response({"questions": QuizQuestionManageSerializer(questions, many=True).data})

        completed = TrainingProgress.objects.filter(
            user=request.user, training=training, status=TrainingProgress.STATUS_COMPLETED,
        ).exists()
        if not completed:
            raise PermissionDenied("Quiz video sona qədər izlənildikdən sonra açılır.")
        return Response({
            "pass_percent": training.pass_percent,
            "questions": QuizQuestionPublicSerializer(questions, many=True).data,
        })

    @action(detail=True, methods=["post"], url_path="quiz/submit")
    def quiz_submit(self, request, pk=None):
        training = self._get_active_training()
        completed = TrainingProgress.objects.filter(
            user=request.user, training=training, status=TrainingProgress.STATUS_COMPLETED,
        ).exists()
        if not completed:
            raise PermissionDenied("Quiz video sona qədər izlənildikdən sonra açılır.")

        questions = list(training.quiz_questions.prefetch_related("options"))
        if not questions:
            raise ValidationError("Bu təlim üçün quiz yoxdur.")

        serializer = QuizSubmitSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        answers = {str(k): v for k, v in serializer.validated_data["answers"].items()}

        missing = [q for q in questions if str(q.id) not in answers]
        if missing:
            raise ValidationError("Bütün suallar cavablandırılmalıdır.")

        with transaction.atomic():
            attempt = QuizAttempt.objects.create(
                user=request.user, training=training, total_count=len(questions),
            )
            correct = 0
            rows = []
            for q in questions:
                options = list(q.options.all())
                selected = next((o for o in options if o.id == answers[str(q.id)]), None)
                if selected is None:
                    raise ValidationError(f"«{q.text[:50]}» sualı üçün yanlış cavab variantı.")
                right = next((o for o in options if o.is_correct), None)
                is_correct = bool(selected.is_correct)
                correct += int(is_correct)
                rows.append(QuizAnswer(
                    attempt=attempt, question=q, selected_option=selected,
                    question_text=q.text, selected_text=selected.text,
                    correct_text=right.text if right else "", is_correct=is_correct,
                ))
            QuizAnswer.objects.bulk_create(rows)
            attempt.correct_count = correct
            attempt.percent = round(correct * 100 / len(questions))
            attempt.passed = attempt.percent >= training.pass_percent
            attempt.save()

        logger.info(f"{request.user} quiz təsdiqlədi: {training} - {attempt.percent}%")
        data = QuizAttemptSerializer(attempt).data
        # İstifadəçiyə yalnız hansı sualın düzgün/səhv olduğu göstərilir, düzgün cavab yox.
        data["results"] = [{"question_id": r.question_id, "is_correct": r.is_correct} for r in rows]
        return Response(data, status=status.HTTP_201_CREATED)

    # ---- Rəy -----------------------------------------------------------------

    @action(detail=True, methods=["get", "post"])
    def feedback(self, request, pk=None):
        training = self.get_object()
        instance = TrainingFeedback.objects.filter(user=request.user, training=training).first()
        if request.method == "GET":
            return Response(TrainingFeedbackSerializer(instance).data if instance else None)

        serializer = TrainingFeedbackSerializer(instance, data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save(user=request.user, training=training)
        return Response(serializer.data, status=status.HTTP_200_OK if instance else status.HTTP_201_CREATED)


class TrainingPermissionsView(APIView):
    """Frontend üçün: cari istifadəçinin Təlimlər modulundakı səlahiyyətləri."""

    module_code = TRAININGS_MODULE_CODE
    permission_classes = [HasSystemAccess, ModuleAccessPermission]

    def get(self, request, *args, **kwargs):
        user = request.user
        return Response({
            "can_view_materials": can_view_materials(user),
            "can_manage_materials": can_manage_materials(user),
            "can_view_statistics": can_view_statistics(user),
            "can_manage_statistics": can_manage_statistics(user),
        })


# --------------------------------------------------------------------------- #
#  Təlim statistikası                                                         #
# --------------------------------------------------------------------------- #

class _StatsScopeMixin:
    permission_classes = [TrainingStatisticsPermission]
    pagination_class = StatsPagination

    def scope(self, qs):
        user = self.request.user
        qs = qs.filter(org_scope_q(user, "training__organization"))
        if not user.is_superuser:
            qs = qs.filter(user__in=statistics_user_scope(user))
        training_id = self.request.query_params.get("training")
        if training_id:
            qs = qs.filter(training_id=training_id)
        return qs


class StatisticsSummaryView(_StatsScopeMixin, APIView):
    """Ümumi göstəricilər və hər təlim üzrə qısa statistika."""

    def get(self, request, *args, **kwargs):
        user = request.user
        trainings = list(
            Training.objects.filter(org_scope_q(user)).order_by("-created_at")
        )
        scoped_user_ids = None if user.is_superuser else set(
            statistics_user_scope(user).values_list("id", flat=True)
        )

        progress_qs = self.scope(TrainingProgress.objects.all())
        attempts_qs = self.scope(QuizAttempt.objects.all())
        feedback_qs = self.scope(TrainingFeedback.objects.all())

        progress_by_training = {}
        for row in progress_qs.values("training_id", "status", "user_id"):
            progress_by_training.setdefault(row["training_id"], []).append(row)

        attempt_stats = {
            r["training_id"]: r for r in attempts_qs.values("training_id").annotate(
                attempts=Count("id"),
                users=Count("user", distinct=True),
                passed_users=Count("user", filter=Q(passed=True), distinct=True),
                avg_percent=Avg("percent"),
            )
        }
        feedback_stats = {
            r["training_id"]: r for r in feedback_qs.values("training_id").annotate(
                count=Count("id"), avg_rating=Avg("rating"),
            )
        }

        rows = []
        totals = {"audience": 0, "completed": 0, "in_progress": 0, "not_started": 0,
                  "quiz_attempts": 0, "feedback": 0}
        for t in trainings:
            records = progress_by_training.get(t.id, [])
            audience = materials_audience_user_ids(t)
            if scoped_user_ids is not None:
                audience &= scoped_user_ids
            started_ids = {r["user_id"] for r in records}
            # Təlimə real baxmış hər kəs (məs. superuser) auditoriyaya daxil sayılır -
            # əks halda "tam baxış" auditoriyadan çox çıxır (133% kimi).
            audience |= started_ids
            completed = sum(1 for r in records if r["status"] == TrainingProgress.STATUS_COMPLETED)
            in_progress = len(records) - completed
            not_started = len(audience - started_ids)
            a = attempt_stats.get(t.id, {})
            f = feedback_stats.get(t.id, {})
            row = {
                "id": t.id,
                "title": t.title,
                "is_active": t.is_active,
                "created_at": t.created_at,
                "audience": len(audience),
                "completed": completed,
                "in_progress": in_progress,
                "not_started": not_started,
                "quiz_attempts": a.get("attempts", 0),
                "quiz_users": a.get("users", 0),
                "quiz_passed_users": a.get("passed_users", 0),
                "quiz_avg_percent": round(a["avg_percent"], 1) if a.get("avg_percent") is not None else None,
                "feedback_count": f.get("count", 0),
                "avg_rating": round(f["avg_rating"], 2) if f.get("avg_rating") is not None else None,
            }
            rows.append(row)
            for key in ("audience", "completed", "in_progress", "not_started"):
                totals[key] += row[key]
            totals["quiz_attempts"] += row["quiz_attempts"]
            totals["feedback"] += row["feedback_count"]

        totals["trainings"] = len(trainings)
        overall_avg = attempts_qs.aggregate(v=Avg("percent"))["v"]
        totals["quiz_avg_percent"] = round(overall_avg, 1) if overall_avg is not None else None
        return Response({
            "totals": totals,
            "trainings": rows,
            "can_manage": can_manage_statistics(user),
        })


class ProgressStatsViewSet(_StatsScopeMixin, mixins.ListModelMixin, mixins.DestroyModelMixin,
                           viewsets.GenericViewSet):
    """
    Kim hansı videoya nə vaxt baxıb. `status`: completed / in_progress.
    DELETE - statistika admini istifadəçinin baxış qeydini sıfırlayır
    (istifadəçi videonu yenidən izləməli olur).
    """

    serializer_class = TrainingProgressStatSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = USER_SEARCH_FIELDS
    ordering_fields = ["completed_at", "first_started_at", "last_heartbeat_at", "attempts_count"]
    ordering = ["-last_heartbeat_at"]

    def get_queryset(self):
        qs = self.scope(
            TrainingProgress.objects.select_related(
                "user", "user__department", "user__organization", "training",
            )
        )
        status_value = self.request.query_params.get("status")
        if status_value in (TrainingProgress.STATUS_COMPLETED, TrainingProgress.STATUS_IN_PROGRESS):
            qs = qs.filter(status=status_value)
        date_field = "completed_at" if status_value == TrainingProgress.STATUS_COMPLETED else "first_started_at"
        return _date_range_filter(qs, self.request, date_field)

    def perform_destroy(self, instance):
        logger.info(f"{self.request.user} baxış qeydini sıfırladı: {instance}")
        instance.delete()

    @action(detail=False, methods=["get"], url_path="not-started")
    def not_started(self, request):
        """Seçilmiş təlimə HƏLƏ baxmağa başlamamış (amma baxmalı olan) istifadəçilər."""
        training_id = request.query_params.get("training")
        training = Training.objects.filter(org_scope_q(request.user)).filter(id=training_id).first()
        if not training:
            raise ValidationError({"training": "Təlim seçilməlidir."})
        audience = materials_audience_user_ids(training)
        if not request.user.is_superuser:
            audience &= set(statistics_user_scope(request.user).values_list("id", flat=True))
        started = set(training.progress_records.values_list("user_id", flat=True))
        users = (
            User.objects.filter(id__in=audience - started)
            .select_related("department", "organization")
            .order_by("firstname", "lastname")
        )
        search = (request.query_params.get("search") or "").strip()
        if search:
            users = users.filter(
                Q(firstname__icontains=search) | Q(lastname__icontains=search)
                | Q(username__icontains=search) | Q(department__title__icontains=search)
            )
        return Response([_user_info(u) for u in users])


class QuizAttemptStatsViewSet(_StatsScopeMixin, mixins.ListModelMixin, mixins.RetrieveModelMixin,
                              mixins.DestroyModelMixin, viewsets.GenericViewSet):
    """Təsdiqlənmiş quiz nəticələri. Detal - hər sualın cavabı ilə."""

    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = USER_SEARCH_FIELDS
    ordering_fields = ["submitted_at", "percent"]
    ordering = ["-submitted_at"]

    def get_serializer_class(self):
        return QuizAttemptDetailSerializer if self.action == "retrieve" else QuizAttemptSerializer

    def get_queryset(self):
        qs = self.scope(
            QuizAttempt.objects.select_related(
                "user", "user__department", "user__organization", "training",
            ).prefetch_related("answers")
        )
        passed = self.request.query_params.get("passed")
        if passed in ("true", "false"):
            qs = qs.filter(passed=(passed == "true"))
        return _date_range_filter(qs, self.request, "submitted_at")

    def perform_destroy(self, instance):
        logger.info(f"{self.request.user} quiz nəticəsini sildi: {instance}")
        instance.delete()


class FeedbackStatsViewSet(_StatsScopeMixin, mixins.ListModelMixin, mixins.DestroyModelMixin,
                           viewsets.GenericViewSet):
    """İstifadəçi rəyləri."""

    serializer_class = TrainingFeedbackStatSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = USER_SEARCH_FIELDS + ["comment"]
    ordering_fields = ["updated_at", "rating"]
    ordering = ["-updated_at"]

    def get_queryset(self):
        qs = self.scope(
            TrainingFeedback.objects.select_related(
                "user", "user__department", "user__organization", "training",
            )
        )
        rating = self.request.query_params.get("rating")
        if rating and rating.isdigit():
            qs = qs.filter(rating=int(rating))
        return _date_range_filter(qs, self.request, "updated_at")


class StatisticsExportView(_StatsScopeMixin, APIView):
    """Statistikanı Excel faylı kimi yükləyir (Baxışlar / Quiz nəticələri / Rəylər)."""

    def get(self, request, *args, **kwargs):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        def fmt(dt):
            return timezone.localtime(dt).strftime("%d.%m.%Y %H:%M") if dt else ""

        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="020624")

        def write_sheet(ws, title, headers, rows):
            ws.title = title
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
            for r in rows:
                ws.append(r)
            for idx, header in enumerate(headers, start=1):
                width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows] or [10])
                ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 12), 60)
            ws.freeze_panes = "A2"

        progress = _date_range_filter(
            self.scope(TrainingProgress.objects.select_related(
                "user", "user__department", "user__organization", "training")),
            request, "first_started_at",
        ).order_by("training__title", "user__firstname")
        write_sheet(wb.active, "Baxışlar",
                    ["Təlim", "Əməkdaş", "Şöbə", "Qurum", "Status", "Başlama sayı",
                     "İlk baxış", "Tam baxılma tarixi", "Son aktivlik"],
                    [[p.training.title, _user_info(p.user)["name"], _user_info(p.user)["department"] or "",
                      _user_info(p.user)["organization"] or "", p.get_status_display(), p.attempts_count,
                      fmt(p.first_started_at), fmt(p.completed_at), fmt(p.last_heartbeat_at)]
                     for p in progress])

        attempts = _date_range_filter(
            self.scope(QuizAttempt.objects.select_related(
                "user", "user__department", "user__organization", "training")),
            request, "submitted_at",
        ).order_by("training__title", "-submitted_at")
        write_sheet(wb.create_sheet(), "Quiz nəticələri",
                    ["Təlim", "Əməkdaş", "Şöbə", "Düzgün", "Sual sayı", "Nəticə (%)", "Keçdi", "Tarix"],
                    [[a.training.title, _user_info(a.user)["name"], _user_info(a.user)["department"] or "",
                      a.correct_count, a.total_count, a.percent, "Bəli" if a.passed else "Xeyr",
                      fmt(a.submitted_at)]
                     for a in attempts])

        feedback = _date_range_filter(
            self.scope(TrainingFeedback.objects.select_related(
                "user", "user__department", "training")),
            request, "updated_at",
        ).order_by("training__title", "-updated_at")
        write_sheet(wb.create_sheet(), "Rəylər",
                    ["Təlim", "Əməkdaş", "Şöbə", "Qiymət", "Rəy", "Tarix"],
                    [[f.training.title, _user_info(f.user)["name"], _user_info(f.user)["department"] or "",
                      f.rating or "", f.comment, fmt(f.updated_at)]
                     for f in feedback])

        buffer = BytesIO()
        wb.save(buffer)
        filename = f"telim_statistikasi_{timezone.localdate():%Y%m%d}.xlsx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
